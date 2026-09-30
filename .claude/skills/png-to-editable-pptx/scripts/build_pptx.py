#!/usr/bin/env python3
"""Build an editable .pptx from a layout spec (JSON) written in source-image pixels.

Usage: build_pptx.py spec.json out.pptx

See references/spec.md for the full schema. Every coordinate is in pixels of the
source PNG; the builder scales them onto a 16:9 (or source-ratio) slide.
"""
import copy
import json
import os
import sys

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.opc.package import Part
from pptx.opc.packuri import PackURI
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import icons  # noqa: E402

SHAPES = {
    "rect": MSO_SHAPE.RECTANGLE,
    "roundRect": MSO_SHAPE.ROUNDED_RECTANGLE,
    "ellipse": MSO_SHAPE.OVAL,
    "chevron": MSO_SHAPE.CHEVRON,
    "homePlate": MSO_SHAPE.PENTAGON,  # arrow-tipped box (first step of a process band)
    "triangle": MSO_SHAPE.ISOSCELES_TRIANGLE,
    "diamond": MSO_SHAPE.DIAMOND,
    "hexagon": MSO_SHAPE.HEXAGON,
    "parallelogram": MSO_SHAPE.PARALLELOGRAM,
    "can": MSO_SHAPE.CAN,
    "rightArrow": MSO_SHAPE.RIGHT_ARROW,
    "leftArrow": MSO_SHAPE.LEFT_ARROW,
    "upArrow": MSO_SHAPE.UP_ARROW,
    "downArrow": MSO_SHAPE.DOWN_ARROW,
    "leftRightArrow": MSO_SHAPE.LEFT_RIGHT_ARROW,
    "upDownArrow": MSO_SHAPE.UP_DOWN_ARROW,
    "cloud": MSO_SHAPE.CLOUD,
    "snip1Rect": MSO_SHAPE.SNIP_1_RECTANGLE,
    "round2SameRect": MSO_SHAPE.ROUND_2_SAME_RECTANGLE,
    "flowProcess": MSO_SHAPE.FLOWCHART_PROCESS,
    "flowDecision": MSO_SHAPE.FLOWCHART_DECISION,
    "noSmoking": MSO_SHAPE.NO_SYMBOL,
}
DASH = {"solid": "solid", "dash": "dash", "dot": "sysDot", "sysDash": "sysDash", "lgDash": "lgDash", "dashDot": "dashDot"}
ALIGN = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT, "j": PP_ALIGN.JUSTIFY}
VALIGN = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}
SVG_EXT_URI = "{96DAC541-7B7A-43D3-8B79-37D633B846F1}"
ASVG_NS = "http://schemas.microsoft.com/office/drawing/2016/SVG/main"


class Builder:
    def __init__(self, spec, base_dir):
        self.spec = spec
        self.base = base_dir
        W, H = spec["source_size"]
        self.W, self.H = W, H
        slide_w_in = spec.get("slide_width_in", 13.333)
        self.emu_per_px = int(slide_w_in * 914400) / W
        self.pt_per_px = slide_w_in * 72 / W
        self.prs = Presentation()
        self.prs.slide_width = Emu(round(W * self.emu_per_px))
        self.prs.slide_height = Emu(round(H * self.emu_per_px))
        self.font = spec.get("font", "Meiryo")
        self.font_latin = spec.get("font_latin", self.font)
        self._svg_n = 0

    # ---- units -----------------------------------------------------------
    def e(self, px):
        return Emu(round(px * self.emu_per_px))

    def pt(self, px):
        return Pt(round(px * self.pt_per_px * 2) / 2)

    @staticmethod
    def rgb(hexstr):
        return RGBColor.from_string(hexstr.lstrip("#").upper()[:6])

    # ---- styling helpers -------------------------------------------------
    def _fill(self, shape, fill):
        if fill is None or fill == "none":
            shape.fill.background()
        elif isinstance(fill, dict):  # {"gradient": [["#hex",0],["#hex",100]], "angle": 90}
            shape.fill.gradient()
            shape.fill.gradient_angle = fill.get("angle", 0)
            stops = fill["gradient"]
            gs_lst = shape.fill._fill._gradFill.gsLst  # noqa: SLF001
            for gs in list(gs_lst):
                gs_lst.remove(gs)
            for color, pos in stops:
                gs = etree.SubElement(gs_lst, qn("a:gs"), pos=str(int(pos * 1000)))
                clr = etree.SubElement(gs, qn("a:srgbClr"), val=color.lstrip("#").upper())
                if len(color.lstrip("#")) == 8:  # #RRGGBBAA
                    etree.SubElement(clr, qn("a:alpha"), val=str(int(int(color[-2:], 16) / 255 * 100000)))
        else:
            shape.fill.solid()
            shape.fill.fore_color.rgb = self.rgb(fill)
            if len(fill.lstrip("#")) == 8:
                self._alpha(shape.fill._xPr.find(qn("a:solidFill"))[0], fill)  # noqa: SLF001

    def _alpha(self, clr_el, hexstr):
        a = int(hexstr.lstrip("#")[6:8], 16) / 255
        etree.SubElement(clr_el, qn("a:alpha"), val=str(int(a * 100000)))

    def _line(self, shape, line):
        ln = shape.line
        if not line:
            ln.fill.background()
            return
        ln.color.rgb = self.rgb(line.get("color", "#000000"))
        ln.width = self.e(line.get("width_px", 1))
        dash = line.get("dash")
        if dash and dash != "solid":
            lnel = ln._get_or_add_ln()  # noqa: SLF001
            for old in lnel.findall(qn("a:prstDash")):
                lnel.remove(old)
            pd = etree.SubElement(lnel, qn("a:prstDash"), val=DASH.get(dash, dash))
            # prstDash must precede round/bevel/miter/headEnd/tailEnd
            lnel.remove(pd)
            idx = 0
            for i, ch in enumerate(lnel):
                if ch.tag in (qn("a:noFill"), qn("a:solidFill"), qn("a:gradFill"), qn("a:pattFill")):
                    idx = i + 1
            lnel.insert(idx, pd)

    def _arrowheads(self, line_el, head, size="med"):
        if head in (None, "none"):
            return
        for tag, want in (("a:headEnd", head in ("start", "both")), ("a:tailEnd", head in ("end", "both"))):
            if want:
                etree.SubElement(line_el, qn(tag), type="triangle", w=size, len=size)

    def _shadow(self, shape, shadow):
        # drop the theme effect reference so only the explicit effect (or none) applies
        eref = shape._element.find(".//" + qn("a:effectRef"))  # noqa: SLF001
        if eref is not None:
            eref.set("idx", "0")
        if not shadow:
            sp_pr = shape._element.spPr  # noqa: SLF001
            if sp_pr.find(qn("a:effectLst")) is None:
                etree.SubElement(sp_pr, qn("a:effectLst"))
            return
        sp_pr = shape._element.spPr  # noqa: SLF001
        eff = etree.SubElement(sp_pr, qn("a:effectLst"))
        sh = etree.SubElement(
            eff, qn("a:outerShdw"), blurRad=str(int(self.e(shadow.get("blur_px", 6)))),
            dist=str(int(self.e(shadow.get("dist_px", 2)))), dir="5400000", algn="t", rotWithShape="0",
        )
        clr = etree.SubElement(sh, qn("a:srgbClr"), val=shadow.get("color", "#000000").lstrip("#")[:6])
        etree.SubElement(clr, qn("a:alpha"), val=str(int(shadow.get("alpha", 0.25) * 100000)))

    # ---- text --------------------------------------------------------------
    def _runs(self, el):
        """Normalize text spec to list of paragraphs, each a list of run dicts."""
        if "paras" in el:
            return [[r if isinstance(r, dict) else {"text": r} for r in p] if isinstance(p, list) else [{"text": p}] for p in el["paras"]]
        return [[{"text": line}] for line in str(el.get("text", "")).split("\n")]

    def _text(self, tf, el):
        tf.word_wrap = el.get("wrap", False)
        tf.auto_size = MSO_AUTO_SIZE.NONE
        ins = el.get("inset_px", [0, 0, 0, 0])
        if isinstance(ins, (int, float)):
            ins = [ins] * 4
        tf.margin_left, tf.margin_top, tf.margin_right, tf.margin_bottom = (self.e(v) for v in ins)
        tf.vertical_anchor = VALIGN[el.get("valign", "m")]
        body = tf._txBody.find(qn("a:bodyPr"))  # noqa: SLF001
        if el.get("vertical"):
            body.set("vert", "eaVert")
        paras = self._runs(el)
        for i, runs in enumerate(paras):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = ALIGN[el.get("align", "l")]
            if el.get("line_pitch_px"):  # exact baseline-to-baseline distance, font-independent
                p.line_spacing = self.pt(el["line_pitch_px"])
            elif el.get("line_spacing"):
                p.line_spacing = el["line_spacing"]
            for rs in runs:
                r = p.add_run()
                r.text = rs.get("text", "")
                f = r.font
                f.size = self.pt(rs.get("size_px", el.get("size_px", 16)))
                f.bold = rs.get("bold", el.get("bold", False))
                f.italic = rs.get("italic", el.get("italic", False))
                f.color.rgb = self.rgb(rs.get("color", el.get("color", "#000000")))
                font = rs.get("font", el.get("font", self.font))
                f.name = rs.get("font_latin", el.get("font_latin", self.font_latin if font == self.font else font))
                rpr = r._r.get_or_add_rPr()  # noqa: SLF001
                for tag in ("a:ea", "a:cs"):
                    etree.SubElement(rpr, qn(tag), typeface=font)

    # ---- element builders ------------------------------------------------
    def add_shape(self, slide, el):
        x, y, w, h = el["box"]
        kind = el.get("type", "rect")
        shp = slide.shapes.add_shape(SHAPES[kind], self.e(x), self.e(y), self.e(w), self.e(h))
        if kind == "roundRect" and "radius_px" in el:
            shp.adjustments[0] = max(0.0, min(0.5, el["radius_px"] / min(w, h)))
        elif kind in ("chevron", "homePlate") and "point_px" in el:
            shp.adjustments[0] = max(0.0, min(1.0, el["point_px"] / h))
        elif "adj" in el:
            for i, v in enumerate(el["adj"]):
                shp.adjustments[i] = v
        if el.get("flip_h"):
            shp._element.spPr.find(qn("a:xfrm")).set("flipH", "1")  # noqa: SLF001
        if el.get("rotation"):
            shp.rotation = el["rotation"]
        self._fill(shp, el.get("fill", "#FFFFFF"))
        self._line(shp, el.get("line"))
        self._shadow(shp, el.get("shadow"))
        if "text" in el or "paras" in el:
            self._text(shp.text_frame, {"align": "c", **el})
        return shp

    def add_text(self, slide, el):
        x, y, w, h = el["box"]
        tb = slide.shapes.add_textbox(self.e(x), self.e(y), self.e(w), self.e(h))
        if el.get("fill"):
            self._fill(tb, el["fill"])
        self._text(tb.text_frame, el)
        return tb

    def add_line(self, slide, el):
        pts = el["points"]
        head = el.get("head", "end" if el["type"] == "arrow" else "none")
        if len(pts) == 2 and not el.get("elbow"):
            (x1, y1), (x2, y2) = pts
            shp = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, self.e(x1), self.e(y1), self.e(x2), self.e(y2))
        else:
            if el.get("elbow") and len(pts) == 2:
                (x1, y1), (x2, y2) = pts
                pts = [(x1, y1), (x1, y2), (x2, y2)] if el["elbow"] == "v" else [(x1, y1), (x2, y1), (x2, y2)]
            fb = slide.shapes.build_freeform(self.e(pts[0][0]), self.e(pts[0][1]), scale=1.0)
            fb.add_line_segments([(self.e(px), self.e(py)) for px, py in pts[1:]], close=False)
            shp = fb.convert_to_shape()
            shp.fill.background()
        self._line(shp, {"color": el.get("color", "#000000"), "width_px": el.get("width_px", 2), "dash": el.get("dash")})
        self._arrowheads(shp.line._get_or_add_ln(), head, el.get("head_size", "med"))  # noqa: SLF001
        return shp

    def _add_svg_picture(self, slide, svg_text, png_bytes, x, y, w, h, name):
        import io

        pic = slide.shapes.add_picture(io.BytesIO(png_bytes), self.e(x), self.e(y), self.e(w), self.e(h))
        pic.name = name
        self._svg_n += 1
        part = Part(PackURI(f"/ppt/media/icon_svg{self._svg_n}.svg"), "image/svg+xml", self.prs.part.package, svg_text.encode("utf-8"))
        rid = slide.part.relate_to(part, RT.IMAGE)
        blip = pic._element.find(".//" + qn("a:blip"))  # noqa: SLF001
        ext_lst = etree.SubElement(blip, qn("a:extLst"))
        ext = etree.SubElement(ext_lst, qn("a:ext"), uri=SVG_EXT_URI)
        svg_blip = etree.SubElement(ext, f"{{{ASVG_NS}}}svgBlip", nsmap={"asvg": ASVG_NS})
        svg_blip.set(qn("r:embed"), rid)
        return pic

    def add_icon(self, slide, el):
        lib, name = icons.parse_ref(el["icon"])
        x, y, w, h = el["box"]
        svg = icons.get_svg(lib, name, color=el.get("color"), stroke_width=el.get("stroke_width"))
        side = max(64, int(max(w, h) * 4))
        png = icons.render_png(svg, side)
        return self._add_svg_picture(slide, svg, png, x, y, w, h, f"icon {lib}:{name}")

    def add_image(self, slide, el):
        x, y, w, h = el["box"]
        path = el["path"] if os.path.isabs(el["path"]) else os.path.join(self.base, el["path"])
        return slide.shapes.add_picture(path, self.e(x), self.e(y), self.e(w), self.e(h))

    def add_table(self, slide, el):
        x, y, w, h = el["box"]
        rows = el["cells"]
        nr, nc = len(rows), max(len(r) for r in rows)
        gf = slide.shapes.add_table(nr, nc, self.e(x), self.e(y), self.e(w), self.e(h))
        tbl = gf.table
        tbl_pr = tbl._tbl.tblPr  # noqa: SLF001
        for attr in ("firstRow", "bandRow"):
            tbl_pr.set(attr, "0")
        style = tbl_pr.find(qn("a:tableStyleId"))
        if style is not None:
            style.text = "{5940675A-B579-460E-94D1-54222C63F5DA}"  # "No Style, Table Grid"
        col_w = el.get("col_widths_px") or [w / nc] * nc
        row_h = el.get("row_heights_px") or [h / nr] * nr
        for i, cw in enumerate(col_w):
            tbl.columns[i].width = self.e(cw)
        for i, rh in enumerate(row_h):
            tbl.rows[i].height = self.e(rh)
        base = {"size_px": el.get("size_px", 14), "color": el.get("color", "#000000"), "align": el.get("align", "c"), "valign": "m", "wrap": True}
        border = el.get("border", {"color": "#999999", "width_px": 1})
        for r, row in enumerate(rows):
            for c in range(nc):
                cell_spec = row[c] if c < len(row) else ""
                if cell_spec is None:
                    continue
                if not isinstance(cell_spec, dict):
                    cell_spec = {"text": str(cell_spec)}
                cell = tbl.cell(r, c)
                if cell_spec.get("span"):
                    rs, cs = cell_spec["span"]
                    cell.merge(tbl.cell(r + rs - 1, c + cs - 1))
                fill = cell_spec.get("fill", (el.get("header_fill") if r == 0 else None) or el.get("fill"))
                if fill:
                    cell.fill.solid()
                    cell.fill.fore_color.rgb = self.rgb(fill)
                else:
                    cell.fill.background()
                sp = {**base, **({"bold": True, "color": el.get("header_color", base["color"])} if r == 0 and el.get("header_fill") else {}), **cell_spec}
                pad = el.get("cell_pad_px", 4)
                ins = sp.get("inset_px", [pad, 2, pad, 2])
                if isinstance(ins, (int, float)):
                    ins = [ins] * 4
                self._text(cell.text_frame, sp)
                # table cells ignore bodyPr insets; padding lives on the cell (tcPr marL/marT/...)
                cell.margin_left, cell.margin_top, cell.margin_right, cell.margin_bottom = (self.e(v) for v in ins)
                cell.vertical_anchor = VALIGN[sp.get("valign", "m")]
                tc_pr = cell._tc.get_or_add_tcPr()  # noqa: SLF001
                for side in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
                    ln = etree.Element(qn(side), w=str(int(self.e(border["width_px"]))))
                    sf = etree.SubElement(ln, qn("a:solidFill"))
                    etree.SubElement(sf, qn("a:srgbClr"), val=border["color"].lstrip("#"))
                    tc_pr.insert(0, ln)
        return gf

    def add_group(self, slide, el):
        grp = slide.shapes.add_group_shape()
        tmp = _ShapeSink(grp.shapes)
        for child in el["children"]:
            self.add(tmp, child)
        grp.name = el.get("name", "group")
        return grp

    def add(self, slide, el):
        t = el.get("type", "rect")
        if t == "text":
            s = self.add_text(slide, el)
        elif t in ("line", "arrow"):
            s = self.add_line(slide, el)
        elif t == "icon":
            s = self.add_icon(slide, el)
        elif t == "image":
            s = self.add_image(slide, el)
        elif t == "table":
            s = self.add_table(slide, el)
        elif t == "group":
            s = self.add_group(slide, el)
        elif t in SHAPES:
            s = self.add_shape(slide, el)
        else:
            raise ValueError(f"unknown element type {t!r}")
        if el.get("name") and t != "group":
            s.name = el["name"]
        return s

    def build(self, out):
        slides = self.spec.get("slides") or [{"elements": self.spec["elements"], "background": self.spec.get("background")}]
        for sd in slides:
            slide = self.prs.slides.add_slide(self.prs.slide_layouts[6])
            bg = sd.get("background", "#FFFFFF")
            if bg:
                slide.background.fill.solid()
                slide.background.fill.fore_color.rgb = self.rgb(bg)
            for el in sd["elements"]:
                self.add(_ShapeSink(slide.shapes, slide), el)
            # No theme effects anywhere (connectors/freeforms/pictures included): shadows exist
            # only where a spec element asks for one via an explicit effectLst.
            for eref in slide._element.iter(qn("a:effectRef")):  # noqa: SLF001
                eref.set("idx", "0")
            if sd.get("notes"):
                slide.notes_slide.notes_text_frame.text = sd["notes"]
        self.prs.save(out)
        return out


class _ShapeSink:
    """Lets builders call .shapes on either a slide or a group, and find the slide part."""

    def __init__(self, shapes, slide=None):
        self.shapes = shapes
        self._slide = slide

    @property
    def part(self):
        return self.shapes.part


def main():
    if len(sys.argv) != 3:
        print(__doc__)
        return 2
    spec_path, out = sys.argv[1], sys.argv[2]
    spec = json.load(open(spec_path, encoding="utf-8"))
    Builder(spec, os.path.dirname(os.path.abspath(spec_path))).build(out)
    print(out)
    return 0


if __name__ == "__main__":
    sys.exit(main())
